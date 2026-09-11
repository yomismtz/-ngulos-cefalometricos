import csv
import json
import math
import os
import re
import sqlite3
import statistics
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, simpledialog, ttk

import yomceph_desktop as core
import yomceph_desktop_hidpi as ui
import yomceph_desktop_v11_database as v11db
import yomceph_desktop_v116 as v116
import yomceph_desktop_v116_final as v116final
from yomceph_scientific_catalog import (
    ANALYSES,
    EXTRA_POINTS,
    MEASUREMENTS,
    STATUS_ACTIVE,
    STATUS_OTHER_PROJECTION,
    STATUS_REFERENCE_ONLY,
    STATUS_VALIDATION,
    analysis_measurements,
    required_landmarks,
)

APP_VERSION = "0.12.0"
MAX_RESEARCH_CASES = 1000
UAM_STUDY_ID = "uam2026"
UAM_CLINICS = ["Tepepan", "Tláhuac", "San Lorenzo", "Nezahualcóyotl"]

CLASSIC_STEINER = {
    "SNA": (82.0, 2.0, "°"),
    "SNB": (80.0, 2.0, "°"),
    "ANB": (2.0, 2.0, "°"),
    "SN–GoGn": (32.0, 5.0, "°"),
    "Plano oclusal–SN": (14.0, 2.0, "°"),
    "IS–NA (°)": (22.0, 2.0, "°"),
    "IS–NA (mm)": (4.0, 2.0, "mm"),
    "II–NB (°)": (25.0, 2.0, "°"),
    "II–NB (mm)": (4.0, 1.0, "mm"),
    "Línea S · labio superior": (0.0, 2.0, "mm"),
    "Línea S · labio inferior": (0.0, 2.0, "mm"),
}

POWELL_RANGES = {
    "Powell Nasofrontal": (115.0, 130.0),
    "Powell Nasofacial": (30.0, 40.0),
    "Powell Nasomental": (120.0, 132.0),
    "Powell Mentocervical": (80.0, 95.0),
}

JARABAK_REFERENCE = {
    "Jarabak · Saddle": (123.0, 6.0, "°"),
    "Jarabak · Articular": (143.0, 5.0, "°"),
    "Jarabak · Gonial": (130.0, 6.0, "°"),
    "Jarabak · Sum": (396.0, 5.0, "°"),
    "Jarabak · S-Ar": (32.0, 3.0, "mm"),
    "Jarabak · S-N": (71.0, 3.0, "mm"),
    "Jarabak · Ar-Go": (44.0, 5.0, "mm"),
    "Jarabak · Go-Me": (71.0, 5.0, "mm"),
    "Jarabak · S-Go": (77.5, 7.5, "mm"),
    "Jarabak · N-Me": (112.5, 7.5, "mm"),
    "Jarabak ratio": (63.5, 1.5, "%"),
}


def _safe_id(text):
    value = re.sub(r"[^A-Za-z0-9_-]+", "_", (text or "").strip()).strip("_")
    return value[:60] or datetime.now().strftime("study_%Y%m%d_%H%M%S")


def _loads(raw, default):
    try:
        return json.loads(raw or "")
    except Exception:
        return default


def _project_point(p, a, b):
    dx, dy = b[0] - a[0], b[1] - a[1]
    den = dx * dx + dy * dy
    if den <= 1e-12:
        return a
    t = ((p[0] - a[0]) * dx + (p[1] - a[1]) * dy) / den
    return (a[0] + t * dx, a[1] + t * dy)


def _clean_spss_name(value):
    name = re.sub(r"[^A-Za-z0-9_]+", "_", value or "var").strip("_")
    if not name or name[0].isdigit():
        name = "v_" + name
    return name[:60]


class YomCephV120(v116final.YomCephV116Final):
    """Flujo modular para caso individual e investigación.

    v0.12 mantiene las fórmulas heredadas y añade una capa de protocolo. Los
    análisis nuevos sólo se calculan cuando el catálogo científico los marca
    como activos.
    """

    def __init__(self):
        self.workflow_type = "individual"
        self.current_study_id = None
        self.current_study = None
        self.selected_measure_keys = []
        self.steiner_reference = "none"
        self.case_group_values = {}
        self.manual_eligibility = "pending"
        self._startup_home_shown = False
        self._legacy_top_widgets = []
        self._left_pane = None
        self._center_pane = None
        self._right_pane = None
        self._left_visible = True
        self._right_visible = True
        self._resize_job = None
        self._extend_master_points()
        super().__init__()
        self.title("YomCeph Desktop · v0.12.0")
        self._capture_panes()
        self.bind("<Configure>", self._on_window_resize, add="+")
        self.after(250, self._responsive_layout)
        self.after(350, self._show_workflow_home)

    def _extend_master_points(self):
        seen = {p[0] for p in v11db.MASTER_POINTS}
        for item in EXTRA_POINTS:
            if item[0] not in seen:
                v11db.MASTER_POINTS.append(item)
                seen.add(item[0])
        v11db.MASTER_ORDER[:] = [p[0] for p in v11db.MASTER_POINTS]
        v11db.MASTER_INFO.clear()
        v11db.MASTER_INFO.update({p[0]: (p[1], p[2]) for p in v11db.MASTER_POINTS})

    def _build_ui(self):
        super()._build_ui()
        try:
            slaves = list(self.pack_slaves())
            idx = slaves.index(self.body_panes)
            for widget in slaves[:idx]:
                self._legacy_top_widgets.append(widget)
                widget.pack_forget()
        except Exception:
            pass

    def _install_database_bar(self):
        self._ensure_metadata_vars()
        bar = ttk.Frame(self, padding=(6, 3))
        try:
            bar.pack(fill=tk.X, before=self.body_panes)
        except Exception:
            bar.pack(fill=tk.X)
        self.universal_wrapper = bar
        self.workflow_badge = tk.Label(bar, text="CASO INDIVIDUAL", bg=ui.C["purple_dark"], fg="white", padx=8, pady=3, font=("Segoe UI", 9, "bold"))
        self.workflow_badge.pack(side=tk.LEFT, padx=(0, 6))
        self.study_compact_label = ttk.Label(bar, text="Sin investigación", style="Section.TLabel")
        self.study_compact_label.pack(side=tk.LEFT, padx=(0, 8))
        ttk.Label(bar, text="Caso:", style="Muted.TLabel").pack(side=tk.LEFT)
        ttk.Entry(bar, textvariable=self.case_id, width=9).pack(side=tk.LEFT, padx=(3, 7))
        ttk.Button(bar, text="📂 RX", command=self.open_image, style="Purple.TButton").pack(side=tk.LEFT, padx=2)
        ttk.Button(bar, text="📏 Calibrar", command=self.start_calibration, style="Turquoise.TButton").pack(side=tk.LEFT, padx=2)
        ttk.Button(bar, text="✓ Calcular", command=self.calculate, style="Gold.TButton").pack(side=tk.LEFT, padx=2)
        ttk.Button(bar, text="💾 Guardar", command=self.save_case_to_database, style="Mint.TButton").pack(side=tk.LEFT, padx=2)
        ttk.Button(bar, text="＋ Caso", command=self.next_case, style="Soft.TButton").pack(side=tk.LEFT, padx=2)
        self.db_counter = ttk.Label(bar, text="", style="Muted.TLabel")
        self.db_counter.pack(side=tk.RIGHT, padx=(4, 8))
        more = ttk.Frame(bar); more.pack(side=tk.RIGHT)
        ttk.Button(more, text="Datos", command=self.open_case_data, style="Soft.TButton").pack(side=tk.LEFT, padx=2)
        ttk.Button(more, text="Base", command=self.open_database_browser, style="Soft.TButton").pack(side=tk.LEFT, padx=2)
        ttk.Button(more, text="Exportar", command=self.export_research_package, style="Soft.TButton").pack(side=tk.LEFT, padx=2)
        ttk.Button(more, text="Puntos", command=self.toggle_left_panel, style="Soft.TButton").pack(side=tk.LEFT, padx=2)
        ttk.Button(more, text="Resultados", command=self.toggle_right_panel, style="Soft.TButton").pack(side=tk.LEFT, padx=2)
        self.case_state_frame = tk.Frame(self, bg="#EEF2F6", padx=8, pady=2)
        try:
            self.case_state_frame.pack(fill=tk.X, before=self.body_panes)
        except Exception:
            self.case_state_frame.pack(fill=tk.X)
        self.case_state_label = tk.Label(self.case_state_frame, text="CASO — · NUEVO", bg="#EEF2F6", fg="#233044", font=("Segoe UI", 9, "bold"), anchor="w")
        self.case_state_label.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.qc_button = ttk.Button(self.case_state_frame, text="QC pendiente", command=self.show_quality_control, style="Soft.TButton")
        self.qc_button.pack(side=tk.RIGHT)
        self.profile_combo = self.language_combo = self.interface_combo = None
        self.study_row = self.meta1 = self.meta2 = None
        self.open_db_button = self.export_db_button = self.more_info_button = None

    def _apply_profile_rules(self, initial=False): return
    def _apply_interface_mode(self): return
    def _apply_language(self):
        try: self.title("YomCeph Desktop · v0.12.0")
        except Exception: pass

    def choose_analysis_mode(self, startup=False):
        if not startup:
            if self.workflow_type == "research": self.open_protocol_editor()
            else: self._open_analysis_selector("Caso individual · seleccione resultados", self._finish_individual_setup)

    def _capture_panes(self):
        try:
            panes = list(self.body_panes.panes())
            if len(panes) >= 3:
                self._left_pane = self.nametowidget(panes[0]); self._center_pane = self.nametowidget(panes[1]); self._right_pane = self.nametowidget(panes[2])
                self.body_panes.pane(self._left_pane, weight=1); self.body_panes.pane(self._center_pane, weight=8); self.body_panes.pane(self._right_pane, weight=2)
        except Exception: pass

    def _on_window_resize(self, _event=None):
        if self._resize_job:
            try: self.after_cancel(self._resize_job)
            except Exception: pass
        self._resize_job = self.after(120, self._responsive_layout)

    def _pane_present(self, pane):
        try: return pane is not None and str(pane) in self.body_panes.panes()
        except Exception: return False

    def _show_pane(self, pane, position, weight):
        if pane is None or self._pane_present(pane): return
        try: self.body_panes.insert(position, pane, weight=weight)
        except Exception:
            try: self.body_panes.add(pane, weight=weight)
            except Exception: pass

    def _hide_pane(self, pane):
        if pane is None or not self._pane_present(pane): return
        try: self.body_panes.forget(pane)
        except Exception: pass

    def _responsive_layout(self):
        if not self._center_pane: self._capture_panes()
        w = max(1, self.winfo_width())
        if w < 1000:
            self._hide_pane(self._left_pane); self._hide_pane(self._right_pane); self._left_visible = self._right_visible = False
        elif w < 1320:
            self._show_pane(self._left_pane, 0, 1); self._hide_pane(self._right_pane); self._left_visible, self._right_visible = True, False
        else:
            self._show_pane(self._left_pane, 0, 1); self._show_pane(self._right_pane, "end", 2); self._left_visible = self._right_visible = True
        try:
            if self._center_pane: self.body_panes.pane(self._center_pane, weight=8)
        except Exception: pass

    def toggle_left_panel(self):
        if self._pane_present(self._left_pane): self._hide_pane(self._left_pane); self._left_visible = False
        else: self._show_pane(self._left_pane, 0, 1); self._left_visible = True

    def toggle_right_panel(self):
        if self._pane_present(self._right_pane): self._hide_pane(self._right_pane); self._right_visible = False
        else: self._show_pane(self._right_pane, "end", 2); self._right_visible = True

    def _init_database(self):
        super()._init_database()
        with sqlite3.connect(self._db_path) as con:
            con.execute("""CREATE TABLE IF NOT EXISTS research_studies (study_id TEXT PRIMARY KEY,name TEXT NOT NULL,target_n INTEGER NOT NULL,allow_target_increase INTEGER NOT NULL DEFAULT 1,age_min INTEGER,age_max INTEGER,inclusion_criteria TEXT,exclusion_criteria TEXT,country TEXT,institution TEXT,group_fields_json TEXT NOT NULL DEFAULT '[]',analyses_json TEXT NOT NULL DEFAULT '[]',measurements_json TEXT NOT NULL DEFAULT '[]',steiner_reference TEXT NOT NULL DEFAULT 'none',protocol_version INTEGER NOT NULL DEFAULT 1,protocol_locked INTEGER NOT NULL DEFAULT 0,created_at TEXT NOT NULL,updated_at TEXT NOT NULL)""")
            existing = {row[1] for row in con.execute("PRAGMA table_info(cases)")}
            additions = {"workflow_type":"TEXT","research_study_id":"TEXT","local_case_id":"TEXT","study_groups_json":"TEXT","eligibility_status":"TEXT","eligibility_reason":"TEXT","selected_measurements_json":"TEXT","steiner_reference":"TEXT"}
            for name, sql_type in additions.items():
                if name not in existing: con.execute(f"ALTER TABLE cases ADD COLUMN {name} {sql_type}")
            now = datetime.now().isoformat(timespec="seconds")
            uam_measures = [k for k in list(v11db.STEINER_ORDER) + list(v11db.POSTURE_ORDER) if k in MEASUREMENTS]
            con.execute("""INSERT OR IGNORE INTO research_studies(study_id,name,target_n,allow_target_increase,age_min,age_max,inclusion_criteria,exclusion_criteria,country,institution,group_fields_json,analyses_json,measurements_json,steiner_reference,protocol_version,protocol_locked,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(UAM_STUDY_ID,"Investigación UAM 2026",103,1,0,15,"Según protocolo aprobado de la investigación UAM 2026.","Según protocolo aprobado de la investigación UAM 2026.","México","UAM",json.dumps([{"name":"Clínica","options":UAM_CLINICS}],ensure_ascii=False),json.dumps(["steiner","posture","powell"],ensure_ascii=False),json.dumps(uam_measures,ensure_ascii=False),"uam2026",1,1,now,now))
            con.execute("""UPDATE cases SET workflow_type='research', research_study_id=?, local_case_id=COALESCE(NULLIF(local_case_id,''),case_id), steiner_reference='uam2026' WHERE (app_profile='uam2026' OR study_name='Investigación UAM 2026') AND (workflow_type IS NULL OR workflow_type='')""",(UAM_STUDY_ID,))
            con.execute("""UPDATE cases SET workflow_type='individual', local_case_id=COALESCE(NULLIF(local_case_id,''),case_id) WHERE workflow_type IS NULL OR workflow_type=''""")
            con.commit()

    def _row_to_study(self,row):
        if not row: return None
        d=dict(row); d["group_fields"]=_loads(d.pop("group_fields_json","[]"),[]); d["analyses"]=_loads(d.pop("analyses_json","[]"),[]); d["measurements"]=_loads(d.pop("measurements_json","[]"),[]); return d

    def _get_study(self,study_id):
        if not study_id: return None
        with sqlite3.connect(self._db_path) as con:
            con.row_factory=sqlite3.Row; row=con.execute("SELECT * FROM research_studies WHERE study_id=?",(study_id,)).fetchone()
        return self._row_to_study(row)

    def _list_studies(self):
        with sqlite3.connect(self._db_path) as con: return con.execute("SELECT study_id,name,target_n,protocol_locked FROM research_studies ORDER BY created_at").fetchall()

    def _save_study(self,data,existing_id=None):
        now=datetime.now().isoformat(timespec="seconds"); study_id=existing_id or _safe_id(data["name"])
        with sqlite3.connect(self._db_path) as con:
            if not existing_id:
                base,i=study_id,2
                while con.execute("SELECT 1 FROM research_studies WHERE study_id=?",(study_id,)).fetchone(): study_id=f"{base}_{i}"; i+=1
                con.execute("""INSERT INTO research_studies(study_id,name,target_n,allow_target_increase,age_min,age_max,inclusion_criteria,exclusion_criteria,country,institution,group_fields_json,analyses_json,measurements_json,steiner_reference,protocol_version,protocol_locked,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(study_id,data["name"],data["target_n"],int(data["allow_target_increase"]),data["age_min"],data["age_max"],data["inclusion"],data["exclusion"],data["country"],data["institution"],json.dumps(data["group_fields"],ensure_ascii=False),json.dumps(data["analyses"],ensure_ascii=False),json.dumps(data["measurements"],ensure_ascii=False),data["steiner_reference"],1,0,now,now))
            else:
                con.execute("""UPDATE research_studies SET name=?,target_n=?,allow_target_increase=?,age_min=?,age_max=?,inclusion_criteria=?,exclusion_criteria=?,country=?,institution=?,group_fields_json=?,analyses_json=?,measurements_json=?,steiner_reference=?,updated_at=? WHERE study_id=?""",(data["name"],data["target_n"],int(data["allow_target_increase"]),data["age_min"],data["age_max"],data["inclusion"],data["exclusion"],data["country"],data["institution"],json.dumps(data["group_fields"],ensure_ascii=False),json.dumps(data["analyses"],ensure_ascii=False),json.dumps(data["measurements"],ensure_ascii=False),data["steiner_reference"],now,study_id))
            con.commit()
        return study_id

    def _show_workflow_home(self):
        if self._startup_home_shown: return
        self._startup_home_shown=True; win=tk.Toplevel(self); win.title("YomCeph · ¿Qué vamos a hacer hoy?"); win.transient(self); win.grab_set(); win.configure(bg=ui.C["paper"]); win.protocol("WM_DELETE_WINDOW",lambda:(win.destroy(),self.destroy()))
        sw,sh=self.winfo_screenwidth(),self.winfo_screenheight(); w,h=min(790,int(sw*.9)),min(560,int(sh*.84)); win.geometry(f"{w}x{h}+{max(0,(sw-w)//2)}+{max(0,(sh-h)//2)}")
        head=tk.Frame(win,bg=ui.C["purple"],padx=28,pady=20); head.pack(fill=tk.X); tk.Label(head,text="¿Qué vamos a hacer hoy?",bg=ui.C["purple"],fg="white",font=("Segoe UI Semibold",22)).pack(anchor="w"); tk.Label(head,text="Elige el objetivo. YomCeph mostrará sólo los datos y puntos necesarios.",bg=ui.C["purple"],fg=ui.C["lilac_soft"],font=("Segoe UI",10)).pack(anchor="w",pady=(4,0))
        body=ttk.Frame(win,padding=24); body.pack(fill=tk.BOTH,expand=True); body.columnconfigure(0,weight=1); body.columnconfigure(1,weight=1); body.rowconfigure(0,weight=1)
        ind=ttk.LabelFrame(body,text="👤 Caso individual",padding=18); ind.grid(row=0,column=0,sticky="nsew",padx=(0,8)); ttk.Label(ind,text="Una radiografía o un paciente para revisar. No entra en estadísticas de investigación.",wraplength=300).pack(anchor="w",pady=(0,18)); ttk.Button(ind,text="Trabajar caso individual",command=lambda:self._start_individual(win),style="Purple.TButton").pack(fill=tk.X)
        res=ttk.LabelFrame(body,text="📊 Investigación",padding=18); res.grid(row=0,column=1,sticky="nsew",padx=(8,0)); ttk.Label(res,text="Define muestra, edades, grupos, criterios, análisis y resultados. Cada estudio mantiene su propia base.",wraplength=300).pack(anchor="w",pady=(0,10)); ttk.Button(res,text="Nueva investigación",command=lambda:self._new_research(win),style="Gold.TButton").pack(fill=tk.X,pady=3); ttk.Button(res,text="Abrir investigación existente",command=lambda:self._choose_existing_study(win),style="Soft.TButton").pack(fill=tk.X,pady=3); ttk.Button(res,text="UAM 2026 · 103 casos",command=lambda:self._activate_study(UAM_STUDY_ID,win),style="Mint.TButton").pack(fill=tk.X,pady=3)

    def _start_individual(self,home=None):
        if home: home.destroy()
        self.workflow_type="individual"; self.current_study_id=None; self.current_study=None; self.profile_var.set(v116.PROFILE_GENERAL); self.study_name_var.set(""); self.study_target_var.set(""); self.case_group_values={}; self.steiner_reference="none"; self._open_analysis_selector("Caso individual · seleccione resultados",self._finish_individual_setup)

    def _finish_individual_setup(self,keys,steiner_reference="none"):
        self.selected_measure_keys=list(keys); self.steiner_reference=steiner_reference; self.analysis_mode="Ambos"; self._apply_selected_measurements(); self._update_compact_context(); self.status.config(text="Caso individual · no pertenece a investigación. Abra una radiografía para comenzar.")

    def _new_research(self,home=None):
        if home: home.destroy()
        self.open_protocol_editor(new_study=True)

    def _choose_existing_study(self,home=None):
        rows=self._list_studies(); win=tk.Toplevel(self); win.title("Abrir investigación"); win.transient(self); win.grab_set(); frame=ttk.Frame(win,padding=14); frame.pack(fill=tk.BOTH,expand=True); tree=ttk.Treeview(frame,columns=("name","target","lock"),show="headings",height=max(5,min(12,len(rows)))); tree.heading("name",text="Investigación"); tree.heading("target",text="Muestra"); tree.heading("lock",text="Protocolo"); tree.column("name",width=360); tree.column("target",width=90,anchor="center"); tree.column("lock",width=110,anchor="center")
        for sid,name,target,locked in rows: tree.insert("",tk.END,iid=sid,values=(name,target,"Bloqueado" if locked else "Editable"))
        tree.pack(fill=tk.BOTH,expand=True)
        def open_selected():
            sel=tree.selection()
            if not sel:return
            win.destroy()
            if home:
                try:home.destroy()
                except Exception:pass
            self._activate_study(sel[0])
        ttk.Button(frame,text="Abrir",command=open_selected,style="Purple.TButton").pack(anchor="e",pady=(8,0))

    def _activate_study(self,study_id,close_window=None):
        study=self._get_study(study_id)
        if not study:messagebox.showerror("Investigación","No se pudo abrir el protocolo.");return
        if close_window:
            try:close_window.destroy()
            except Exception:pass
        self.workflow_type="research"; self.current_study_id=study_id; self.current_study=study; self.selected_measure_keys=[k for k in study["measurements"] if k in MEASUREMENTS]; self.steiner_reference=study.get("steiner_reference") or "none"; self.case_group_values={}; self.manual_eligibility="pending"; self.analysis_mode="Ambos"; self.profile_var.set(v116.PROFILE_UAM if study_id==UAM_STUDY_ID else v116.PROFILE_CUSTOM); self.study_name_var.set(study["name"]); self.study_target_var.set(str(study["target_n"])); self.country_var.set(study.get("country") or ""); self.institution_var.set(study.get("institution") or ""); self._apply_selected_measurements(); self._update_compact_context(); self._update_db_counter(); self.status.config(text=f"Investigación activa: {study['name']}. Capture los datos del caso."); self.open_case_data(first_time=True)

    def open_protocol_editor(self,new_study=False):
        existing=None if new_study else self.current_study
        if existing and existing.get("protocol_locked"):self._show_locked_protocol(existing);return
        win=tk.Toplevel(self);win.title("YomCeph · Protocolo de investigación");win.transient(self);win.grab_set();sw,sh=self.winfo_screenwidth(),self.winfo_screenheight();w,h=min(980,int(sw*.94)),min(820,int(sh*.92));win.geometry(f"{w}x{h}+{max(0,(sw-w)//2)}+{max(0,(sh-h)//2)}");outer=ttk.Frame(win);outer.pack(fill=tk.BOTH,expand=True);canvas=tk.Canvas(outer,bg=ui.C["paper"],highlightthickness=0);scroll=ttk.Scrollbar(outer,orient="vertical",command=canvas.yview);content=ttk.Frame(canvas,padding=18);cid=canvas.create_window((0,0),window=content,anchor="nw");canvas.configure(yscrollcommand=scroll.set);canvas.pack(side=tk.LEFT,fill=tk.BOTH,expand=True);scroll.pack(side=tk.RIGHT,fill=tk.Y);content.bind("<Configure>",lambda e:canvas.configure(scrollregion=canvas.bbox("all")));canvas.bind("<Configure>",lambda e:canvas.itemconfigure(cid,width=e.width));ttk.Label(content,text="Diseño de la investigación",style="Section.TLabel",font=("Segoe UI Semibold",17)).pack(anchor="w");ttk.Label(content,text="El protocolo se aplica a todos los casos y queda bloqueado después del primer caso incluido.",style="Muted.TLabel").pack(anchor="w",pady=(2,10))
        basics=ttk.LabelFrame(content,text="1 · Muestra y criterios",padding=12);basics.pack(fill=tk.X,pady=5);name_var=tk.StringVar(value=(existing or {}).get("name",""));target_var=tk.StringVar(value=str((existing or {}).get("target_n",103)));allow_inc=tk.BooleanVar(value=bool((existing or {}).get("allow_target_increase",1)));amin_var=tk.StringVar(value="" if (existing or {}).get("age_min") is None else str(existing["age_min"]));amax_var=tk.StringVar(value="" if (existing or {}).get("age_max") is None else str(existing["age_max"]));country_var=tk.StringVar(value=(existing or {}).get("country","México"));institution_var=tk.StringVar(value=(existing or {}).get("institution",""));steiner_var=tk.StringVar(value=(existing or {}).get("steiner_reference","none"));grid=ttk.Frame(basics);grid.pack(fill=tk.X)
        for c in (1,3):grid.columnconfigure(c,weight=1)
        ttk.Label(grid,text="Nombre:").grid(row=0,column=0,sticky="w",pady=3);ttk.Entry(grid,textvariable=name_var).grid(row=0,column=1,columnspan=3,sticky="ew",padx=4,pady=3);ttk.Label(grid,text="Muestra planeada:").grid(row=1,column=0,sticky="w",pady=3);ttk.Entry(grid,textvariable=target_var,width=8).grid(row=1,column=1,sticky="w",padx=4,pady=3);ttk.Checkbutton(grid,text="Permitir aumentar la muestra posteriormente",variable=allow_inc).grid(row=1,column=2,columnspan=2,sticky="w",pady=3);ttk.Label(grid,text="Edad mínima:").grid(row=2,column=0,sticky="w",pady=3);ttk.Entry(grid,textvariable=amin_var,width=8).grid(row=2,column=1,sticky="w",padx=4,pady=3);ttk.Label(grid,text="Edad máxima:").grid(row=2,column=2,sticky="w",pady=3);ttk.Entry(grid,textvariable=amax_var,width=8).grid(row=2,column=3,sticky="w",padx=4,pady=3);ttk.Label(grid,text="País:").grid(row=3,column=0,sticky="w",pady=3);ttk.Entry(grid,textvariable=country_var).grid(row=3,column=1,sticky="ew",padx=4,pady=3);ttk.Label(grid,text="Institución:").grid(row=3,column=2,sticky="w",pady=3);ttk.Entry(grid,textvariable=institution_var).grid(row=3,column=3,sticky="ew",padx=4,pady=3);ttk.Label(basics,text="Criterios de inclusión:").pack(anchor="w",pady=(6,2));inclusion=tk.Text(basics,height=3,wrap=tk.WORD);inclusion.pack(fill=tk.X);inclusion.insert("1.0",(existing or {}).get("inclusion_criteria",""));ttk.Label(basics,text="Criterios de exclusión:").pack(anchor="w",pady=(6,2));exclusion=tk.Text(basics,height=3,wrap=tk.WORD);exclusion.pack(fill=tk.X);exclusion.insert("1.0",(existing or {}).get("exclusion_criteria",""))
        groups=ttk.LabelFrame(content,text="2 · Procedencia, grupo o momento",padding=12);groups.pack(fill=tk.X,pady=5);ttk.Label(groups,text="Una variable por línea: Nombre=opción1|opción2|opción3",style="Muted.TLabel").pack(anchor="w");ttk.Label(groups,text="Ej.: Clínica=Tepepan|Tláhuac|San Lorenzo|Nezahualcóyotl · Grupo=Caso|Control · Momento=Pre|Post").pack(anchor="w",pady=(2,5));group_text=tk.Text(groups,height=5,wrap=tk.WORD);group_text.pack(fill=tk.X)
        if existing:group_text.insert("1.0","\n".join(f"{f.get('name','Grupo')}="+"|".join(f.get('options',[])) for f in existing.get("group_fields",[])))
        analyses_box=ttk.LabelFrame(content,text="3 · Análisis y resultados",padding=12);analyses_box.pack(fill=tk.X,pady=5);ttk.Label(analyses_box,text="Puede combinar análisis. Sólo las mediciones con geometría verificada se pueden activar.",style="Muted.TLabel").pack(anchor="w",pady=(0,6));selected=set((existing or {}).get("measurements",[]));analysis_vars={};measure_vars={}
        for aid,spec in ANALYSES.items():
            fr=ttk.LabelFrame(analyses_box,text=spec["name"],padding=7);fr.pack(fill=tk.X,pady=3);active=analysis_measurements(aid,active_only=True);av=tk.BooleanVar(value=any(k in selected for k in active));analysis_vars[aid]=av;cb=ttk.Checkbutton(fr,text="Incluir análisis",variable=av);cb.pack(anchor="w")
            if not active:cb.state(["disabled"])
            status={STATUS_ACTIVE:"Disponible",STATUS_VALIDATION:"En validación",STATUS_REFERENCE_ONLY:"Referencia metodológica",STATUS_OTHER_PROJECTION:"Requiere otra proyección"}.get(spec.get("status"),spec.get("status"));ttk.Label(fr,text=f"{status}. {spec.get('summary','')}",style="Muted.TLabel",wraplength=820,justify=tk.LEFT).pack(anchor="w",padx=(22,0))
            for key in active:
                mv=tk.BooleanVar(value=key in selected);measure_vars[key]=mv;ttk.Checkbutton(fr,text=MEASUREMENTS[key]["label"],variable=mv).pack(anchor="w",padx=(35,0))
        ref=ttk.LabelFrame(content,text="4 · Referencia Steiner",padding=10);ref.pack(fill=tk.X,pady=5);ttk.Radiobutton(ref,text="Sin clasificación automática",value="none",variable=steiner_var).pack(anchor="w");ttk.Radiobutton(ref,text="Steiner clásica publicada",value="classic",variable=steiner_var).pack(anchor="w");ttk.Radiobutton(ref,text="Protocolo UAM 2026",value="uam2026",variable=steiner_var).pack(anchor="w")
        def save_protocol():
            name=name_var.get().strip()
            if not name:messagebox.showwarning("Protocolo","Escriba el nombre.",parent=win);return
            try:
                target=int(target_var.get());
                if not 1<=target<=MAX_RESEARCH_CASES:raise ValueError
                amin=int(amin_var.get()) if amin_var.get().strip() else None;amax=int(amax_var.get()) if amax_var.get().strip() else None
                if amin is not None and not 0<=amin<=130:raise ValueError
                if amax is not None and not 0<=amax<=130:raise ValueError
                if amin is not None and amax is not None and amin>amax:raise ValueError
            except ValueError:messagebox.showwarning("Protocolo","Revise muestra (1–1000) y rango de edad (0–130).",parent=win);return
            fields=[]
            for raw in group_text.get("1.0","end").splitlines():
                raw=raw.strip()
                if not raw:continue
                if "=" not in raw:messagebox.showwarning("Grupos",f"Formato inválido: {raw}",parent=win);return
                fname,opts=raw.split("=",1);options=[x.strip() for x in opts.split("|") if x.strip()]
                if not fname.strip() or not options:messagebox.showwarning("Grupos",f"Variable incompleta: {raw}",parent=win);return
                fields.append({"name":fname.strip(),"options":options})
            measures=[k for k,v in measure_vars.items() if v.get()]
            for aid,v in analysis_vars.items():
                if v.get() and not any(MEASUREMENTS[k]["analysis"]==aid for k in measures):measures.extend(k for k in analysis_measurements(aid,active_only=True) if k not in measures)
            if not measures:messagebox.showwarning("Análisis","Seleccione al menos una medición disponible.",parent=win);return
            aids=[]
            for k in measures:
                aid=MEASUREMENTS[k]["analysis"]
                if aid not in aids:aids.append(aid)
            data=dict(name=name,target_n=target,allow_target_increase=allow_inc.get(),age_min=amin,age_max=amax,inclusion=inclusion.get("1.0","end").strip(),exclusion=exclusion.get("1.0","end").strip(),country=country_var.get().strip(),institution=institution_var.get().strip(),group_fields=fields,analyses=aids,measurements=measures,steiner_reference=steiner_var.get());sid=self._save_study(data,(existing or {}).get("study_id"));win.destroy();self._activate_study(sid)
        buttons=ttk.Frame(content);buttons.pack(fill=tk.X,pady=12);ttk.Button(buttons,text="Guardar protocolo y continuar",command=save_protocol,style="Purple.TButton").pack(side=tk.RIGHT);ttk.Button(buttons,text="Cancelar",command=win.destroy,style="Soft.TButton").pack(side=tk.RIGHT,padx=6)

    def _show_locked_protocol(self,study):
        text=(f"El protocolo '{study['name']}' está bloqueado porque ya contiene casos incluidos.\n\nNo se cambian silenciosamente edades, grupos, análisis ni resultados. ")
        if not study.get("allow_target_increase"):messagebox.showinfo("Protocolo bloqueado",text+f"\nMuestra: {study['target_n']}");return
        if messagebox.askyesno("Protocolo bloqueado",text+f"\nMuestra actual: {study['target_n']}.\n\n¿Aumentar la muestra?"):
            new=simpledialog.askinteger("Muestra","Nuevo número planeado:",minvalue=study["target_n"]+1,maxvalue=MAX_RESEARCH_CASES,parent=self)
            if new:
                with sqlite3.connect(self._db_path) as con:con.execute("UPDATE research_studies SET target_n=?,updated_at=? WHERE study_id=?",(new,datetime.now().isoformat(timespec="seconds"),study["study_id"]));con.commit()
                self._activate_study(study["study_id"])

    def _open_analysis_selector(self,title,on_accept):
        win=tk.Toplevel(self);win.title(title);win.transient(self);win.grab_set();sw,sh=self.winfo_screenwidth(),self.winfo_screenheight();w,h=min(850,int(sw*.92)),min(760,int(sh*.9));win.geometry(f"{w}x{h}+{max(0,(sw-w)//2)}+{max(0,(sh-h)//2)}");outer=ttk.Frame(win);outer.pack(fill=tk.BOTH,expand=True);canvas=tk.Canvas(outer,bg=ui.C["paper"],highlightthickness=0);sc=ttk.Scrollbar(outer,orient="vertical",command=canvas.yview);content=ttk.Frame(canvas,padding=16);cid=canvas.create_window((0,0),window=content,anchor="nw");canvas.configure(yscrollcommand=sc.set);canvas.pack(side=tk.LEFT,fill=tk.BOTH,expand=True);sc.pack(side=tk.RIGHT,fill=tk.Y);content.bind("<Configure>",lambda e:canvas.configure(scrollregion=canvas.bbox("all")));canvas.bind("<Configure>",lambda e:canvas.itemconfigure(cid,width=e.width));ttk.Label(content,text=title,style="Section.TLabel",font=("Segoe UI Semibold",16)).pack(anchor="w",pady=(0,8));vars_={}
        for aid,spec in ANALYSES.items():
            fr=ttk.LabelFrame(content,text=spec["name"],padding=7);fr.pack(fill=tk.X,pady=3);active=analysis_measurements(aid,active_only=True);ttk.Label(fr,text=spec.get("summary",""),style="Muted.TLabel",wraplength=740,justify=tk.LEFT).pack(anchor="w")
            if not active:ttk.Label(fr,text="No calculable todavía: en validación o requiere otra proyección.",style="Muted.TLabel").pack(anchor="w",padx=(20,0))
            for key in active:
                var=tk.BooleanVar(value=key in self.selected_measure_keys);vars_[key]=var;ttk.Checkbutton(fr,text=MEASUREMENTS[key]["label"],variable=var).pack(anchor="w",padx=(20,0))
        ref=tk.StringVar(value=self.steiner_reference or "none");box=ttk.LabelFrame(content,text="Referencia Steiner",padding=8);box.pack(fill=tk.X,pady=5)
        for value,label in [("none","Sin clasificación"),("classic","Steiner clásica"),("uam2026","Protocolo UAM 2026")]:ttk.Radiobutton(box,text=label,value=value,variable=ref).pack(anchor="w")
        def accept():
            keys=[k for k,v in vars_.items() if v.get()]
            if not keys:messagebox.showwarning("Análisis","Seleccione al menos una medición.",parent=win);return
            win.destroy();on_accept(keys,ref.get())
        ttk.Button(content,text="Continuar",command=accept,style="Purple.TButton").pack(anchor="e",pady=10)

    def _sex_code_from_label(self,label):
        for code in v116.SEX_CODES:
            if self._sex_label(code)==label:return code
        return ""

    def _evaluate_eligibility(self):
        if self.workflow_type!="research" or not self.current_study:return "not_applicable","Caso individual"
        st=self.current_study
        try:years=int(self.age_years_var.get().strip())
        except Exception:return "pending","Edad pendiente"
        amin,amax=st.get("age_min"),st.get("age_max")
        if amin is not None and years<amin:return "not_eligible",f"Edad {years} < mínima {amin}"
        if amax is not None and years>amax:return "not_eligible",f"Edad {years} > máxima {amax}"
        if self.manual_eligibility=="no":return "not_eligible","No cumple criterios de inclusión/exclusión"
        if self.manual_eligibility=="yes":return "eligible","Cumple rango de edad y criterios confirmados"
        return "pending","Edad dentro del rango; falta confirmar criterios"

    def open_case_data(self,first_time=False):
        win=tk.Toplevel(self);win.title("Datos del caso");win.transient(self);win.grab_set();sw,sh=self.winfo_screenwidth(),self.winfo_screenheight();w,h=min(720,int(sw*.9)),min(760,int(sh*.9));win.geometry(f"{w}x{h}+{max(0,(sw-w)//2)}+{max(0,(sh-h)//2)}");fr=ttk.Frame(win,padding=16);fr.pack(fill=tk.BOTH,expand=True);ttk.Label(fr,text="Datos del caso",style="Section.TLabel",font=("Segoe UI Semibold",16)).pack(anchor="w")
        if self.workflow_type=="research" and self.current_study:
            st=self.current_study;ttk.Label(fr,text=f"{st['name']} · edad admisible: {st.get('age_min','—')} a {st.get('age_max','—')} años",style="Muted.TLabel").pack(anchor="w",pady=(2,8))
        else:ttk.Label(fr,text="Caso individual · no entra en estadísticas de investigación.",style="Muted.TLabel").pack(anchor="w",pady=(2,8))
        grid=ttk.Frame(fr);grid.pack(fill=tk.X)
        for c in (1,3):grid.columnconfigure(c,weight=1)
        ttk.Label(grid,text="Caso / ID:").grid(row=0,column=0,sticky="w",pady=3);ttk.Entry(grid,textvariable=self.case_id).grid(row=0,column=1,sticky="ew",padx=4,pady=3);ttk.Label(grid,text="Edad (años):").grid(row=1,column=0,sticky="w",pady=3);ttk.Entry(grid,textvariable=self.age_years_var,width=8).grid(row=1,column=1,sticky="w",padx=4,pady=3);ttk.Label(grid,text="Meses:").grid(row=1,column=2,sticky="w",pady=3);ttk.Entry(grid,textvariable=self.age_months_var,width=8).grid(row=1,column=3,sticky="w",padx=4,pady=3);ttk.Label(grid,text="Nacimiento:").grid(row=2,column=0,sticky="w",pady=3);ttk.Entry(grid,textvariable=self.birth_date_var).grid(row=2,column=1,sticky="ew",padx=4,pady=3);ttk.Label(grid,text="Fecha RX:").grid(row=2,column=2,sticky="w",pady=3);ttk.Entry(grid,textvariable=self.radiograph_date_var).grid(row=2,column=3,sticky="ew",padx=4,pady=3);ttk.Label(grid,text="Sexo registrado:").grid(row=3,column=0,sticky="w",pady=3);sex=ttk.Combobox(grid,state="readonly",values=[self._sex_label(c) for c in v116.SEX_CODES]);sex.grid(row=3,column=1,sticky="ew",padx=4,pady=3);sex.set(self._sex_label(self.sex_code_var.get()));ttk.Label(grid,text="Identidad de género (opcional):").grid(row=3,column=2,sticky="w",pady=3);ttk.Entry(grid,textvariable=self.gender_identity_var).grid(row=3,column=3,sticky="ew",padx=4,pady=3)
        group_widgets={}
        if self.workflow_type=="research" and self.current_study:
            gbox=ttk.LabelFrame(fr,text="Procedencia / grupos del estudio",padding=9);gbox.pack(fill=tk.X,pady=8)
            for field in self.current_study.get("group_fields",[]):
                row=ttk.Frame(gbox);row.pack(fill=tk.X,pady=2);ttk.Label(row,text=field.get("name","Grupo")+":",width=22).pack(side=tk.LEFT);var=tk.StringVar(value=self.case_group_values.get(field.get("name"),""));combo=ttk.Combobox(row,textvariable=var,values=field.get("options",[]),state="normal");combo.pack(side=tk.LEFT,fill=tk.X,expand=True);group_widgets[field.get("name")]=var
            crit=ttk.LabelFrame(fr,text="Elegibilidad",padding=9);crit.pack(fill=tk.X,pady=8);ttk.Label(crit,text="YomCeph valida automáticamente la edad. Los demás criterios los confirma el investigador.",style="Muted.TLabel",wraplength=640).pack(anchor="w");elig=tk.StringVar(value=self.manual_eligibility);ttk.Radiobutton(crit,text="Cumple criterios de inclusión/exclusión",value="yes",variable=elig).pack(anchor="w");ttk.Radiobutton(crit,text="No cumple criterios",value="no",variable=elig).pack(anchor="w");ttk.Radiobutton(crit,text="Pendiente de confirmar",value="pending",variable=elig).pack(anchor="w")
        else:elig=tk.StringVar(value="pending")
        note=tk.Label(fr,text="",anchor="w",justify=tk.LEFT,padx=8,pady=5);note.pack(fill=tk.X,pady=(5,0))
        def refresh_note(*_):
            old=self.manual_eligibility;self.manual_eligibility=elig.get();status,reason=self._evaluate_eligibility();self.manual_eligibility=old
            if status=="not_eligible":note.config(text="⚠ NO CORRESPONDE A LA INVESTIGACIÓN · "+reason,bg="#FFF3BF",fg="#6A4B00")
            elif status=="eligible":note.config(text="✓ Elegible para la investigación · "+reason,bg="#E7F7EE",fg="#17633D")
            else:note.config(text="Elegibilidad pendiente · "+reason,bg="#EEF2F6",fg="#344054")
        self.age_years_var.trace_add("write",refresh_note);elig.trace_add("write",refresh_note);refresh_note()
        def accept():
            self.sex_code_var.set(self._sex_code_from_label(sex.get()));self.manual_eligibility=elig.get();self.case_group_values={k:v.get().strip() for k,v in group_widgets.items()}
            if "Clínica" in self.case_group_values:self.clinic_var.set(self.case_group_values["Clínica"])
            status,reason=self._evaluate_eligibility()
            if status=="not_eligible" and self.workflow_type=="research":messagebox.showwarning("Elegibilidad","Este caso no corresponde a la muestra definida. Puede conservarse para trazabilidad, pero no entrará en los cálculos de incluidos.\n\n"+reason,parent=win)
            win.destroy();self._update_compact_context()
        ttk.Button(fr,text="Guardar datos del caso",command=accept,style="Purple.TButton").pack(anchor="e",pady=(8,0))

    def _apply_selected_measurements(self):
        needed=set(required_landmarks(self.selected_measure_keys));self.active_point_order=[k for k in v11db.MASTER_ORDER if k in needed];self._rebuild_active_point_list()
        if hasattr(self,"progress_label"):self.progress_label.config(text=f"{sum(1 for k in self.active_point_order if k in self.points)} / {len(self.active_point_order)} puntos necesarios")

    def _active_measure_order(self):return list(self.selected_measure_keys) if self.selected_measure_keys else super()._active_measure_order()
    def _measurement_unit(self,key):return MEASUREMENTS[key].get("unit","") if key in MEASUREMENTS else super()._measurement_unit(key)

    def calculate_values(self):
        r=super().calculate_values();p=self.points
        if all(k in p for k in ("S","M","G")):
            r["YEN"]=core.angle3(p["S"],p["M"],p["G"]);foot=_project_point(p["M"],p["S"],p["G"]);r["W"]=core.angle3(foot,p["M"],p["G"])
        if self.mm_per_pixel and all(k in p for k in ("A","B","OcA","OcP")):
            ao=_project_point(p["A"],p["OcP"],p["OcA"]);bo=_project_point(p["B"],p["OcP"],p["OcA"]);dx,dy=p["OcA"][0]-p["OcP"][0],p["OcA"][1]-p["OcP"][1];den=math.hypot(dx,dy)
            if den>1e-9:r["Wits AO–BO"]=((ao[0]-bo[0])*dx/den+(ao[1]-bo[1])*dy/den)*self.mm_per_pixel
        if all(k in p for k in ("N","S","Ar")):r["Jarabak · Saddle"]=core.angle3(p["N"],p["S"],p["Ar"])
        if all(k in p for k in ("S","Ar","Go")):r["Jarabak · Articular"]=core.angle3(p["S"],p["Ar"],p["Go"])
        if all(k in p for k in ("Ar","Go","Me")):r["Jarabak · Gonial"]=core.angle3(p["Ar"],p["Go"],p["Me"])
        if all(k in p for k in ("Ar","Go","N")):r["Jarabak · Upper gonial"]=core.angle3(p["Ar"],p["Go"],p["N"])
        if all(k in p for k in ("N","Go","Me")):r["Jarabak · Lower gonial"]=core.angle3(p["N"],p["Go"],p["Me"])
        if all(k in r for k in ("Jarabak · Saddle","Jarabak · Articular","Jarabak · Gonial")):r["Jarabak · Sum"]=r["Jarabak · Saddle"]+r["Jarabak · Articular"]+r["Jarabak · Gonial"]
        for key,a,b in [("Jarabak · S-N","S","N"),("Jarabak · S-Ar","S","Ar"),("Jarabak · Ar-Go","Ar","Go"),("Jarabak · Go-Me","Go","Me"),("Jarabak · S-Go","S","Go"),("Jarabak · N-Me","N","Me")]:
            if self.mm_per_pixel and a in p and b in p:r[key]=core.dist(p[a],p[b])*self.mm_per_pixel
        if all(k in p for k in ("S","Go","N","Me")):
            den=core.dist(p["N"],p["Me"])
            if den>1e-9:r["Jarabak ratio"]=core.dist(p["S"],p["Go"])/den*100.0
        return r

    def _classification_for(self,key,value):
        if key=="YEN":return "Clase II" if value<117 else ("Clase III" if value>123 else "Clase I")
        if key=="W":return "Clase II" if value<51 else ("Clase III" if value>56 else "Clase I")
        if key in POWELL_RANGES:
            lo,hi=POWELL_RANGES[key];return "Bajo" if value<lo else ("Alto" if value>hi else "En rango")
        if key=="Jarabak ratio":return "Tendencia vertical" if value<62 else ("Tendencia horizontal" if value>65 else "Intermedia")
        reference=None
        if self.steiner_reference=="classic":reference=CLASSIC_STEINER.get(key)
        elif self.steiner_reference=="uam2026" and key in v11db.STEINER_PROTOCOL:
            mean,sd,unit,*_=v11db.STEINER_PROTOCOL[key];reference=(mean,sd,unit)
        if reference:
            mean,sd,_=reference;return "Disminuido" if value<mean-sd else ("Aumentado" if value>mean+sd else "En rango")
        return ""

    def _diagnosis_for_result(self,key,value,r):
        cat=self._classification_for(key,value)
        if cat:
            if key in ("YEN","W"):return f"Clasificación según la referencia original del método: {cat}."
            return f"Clasificación según la referencia seleccionada: {cat}. No sustituye diagnóstico clínico."
        if key=="Wits AO–BO":return "Valor sagital AO–BO. La referencia original depende del sexo y del plano oclusal; se conserva el valor crudo."
        if key.startswith("Jarabak"):return "Variable Björk–Jarabak. Interprete con la referencia poblacional/etaria definida por el protocolo."
        try:return super()._diagnosis_for_result(key,value,r)
        except Exception:return ""

    def calculate(self):
        r=self.calculate_values();self.results_cache=r;lines=[]
        for key in self._active_measure_order():
            value=r.get(key)
            if not isinstance(value,(int,float)) or not math.isfinite(value):continue
            unit=self._measurement_unit(key);lines.append(f"{MEASUREMENTS.get(key,{}).get('label',key)}: {value:.2f}{unit}");c=self._classification_for(key,value)
            if c:lines.append(f"Referencia: {c}")
            dx=self._diagnosis_for_result(key,value,r)
            if dx:lines.append(dx)
            lines.append("")
        if not lines:lines=["Aún faltan puntos o calibración para las mediciones seleccionadas."]
        try:self.results_text.config(state=tk.NORMAL);self.results_text.delete("1.0",tk.END);self.results_text.insert("1.0","\n".join(lines));self.results_text.config(state=tk.DISABLED)
        except Exception:pass
        self.status.config(text=f"Análisis calculado · {sum(1 for k in self._active_measure_order() if isinstance(r.get(k),(int,float)))} resultados con valor.")
        try:self.hide_traces()
        except Exception:pass

    def _storage_case_id(self,local_id):
        local=(local_id or "").strip()
        if self.workflow_type=="research" and self.current_study_id:
            if self.current_study_id==UAM_STUDY_ID:
                with sqlite3.connect(self._db_path) as con:
                    if con.execute("SELECT 1 FROM cases WHERE case_id=? AND (research_study_id=? OR app_profile='uam2026')",(local,UAM_STUDY_ID)).fetchone():return local
                    if not con.execute("SELECT 1 FROM cases WHERE case_id=?",(local,)).fetchone():return local
            return f"R:{self.current_study_id}:{local}"
        return f"I:{local}"

    def save_case_to_database(self):
        local=self.case_id.get().strip()
        if not local:messagebox.showwarning("Caso","Escriba el número o identificador del caso.");return
        status,reason=self._evaluate_eligibility()
        if self.workflow_type=="research" and status=="pending" and not messagebox.askyesno("Elegibilidad","La elegibilidad todavía está pendiente. ¿Guardar el caso con estado PENDIENTE?"):return
        storage=self._storage_case_id(local);old_loaded=self._loaded_case_id;self.case_id.set(storage)
        if old_loaded and old_loaded!=storage:self._loaded_case_id=old_loaded
        try:
            if self.workflow_type=="research" and self.current_study:
                self.profile_var.set(v116.PROFILE_UAM if self.current_study_id==UAM_STUDY_ID else v116.PROFILE_CUSTOM);self.study_name_var.set(self.current_study["name"]);self.study_target_var.set(str(self.current_study["target_n"]));self.country_var.set(self.current_study.get("country") or "");self.institution_var.set(self.current_study.get("institution") or "")
                if "Clínica" in self.case_group_values:self.clinic_var.set(self.case_group_values["Clínica"])
            else:self.profile_var.set(v116.PROFILE_GENERAL)
            result=super().save_case_to_database()
            if not self._case_saved:return result
            with sqlite3.connect(self._db_path) as con:
                con.execute("""UPDATE cases SET workflow_type=?,research_study_id=?,local_case_id=?,study_groups_json=?,eligibility_status=?,eligibility_reason=?,selected_measurements_json=?,steiner_reference=? WHERE case_id=?""",(self.workflow_type,self.current_study_id,local,json.dumps(self.case_group_values,ensure_ascii=False),status,reason,json.dumps(self.selected_measure_keys,ensure_ascii=False),self.steiner_reference,storage))
                if self.workflow_type=="research" and status=="eligible" and self.current_study_id:con.execute("UPDATE research_studies SET protocol_locked=1 WHERE study_id=?",(self.current_study_id,))
                con.commit()
            if self.workflow_type=="research" and self.current_study_id:self.current_study=self._get_study(self.current_study_id)
            self.status.config(text=f"✓ Caso {local} guardado · {status.replace('_',' ')}");self._update_db_counter();self._update_compact_context();return result
        finally:self.case_id.set(local)

    def load_case_from_database(self,storage_id):
        with sqlite3.connect(self._db_path) as con:row=con.execute("SELECT workflow_type,research_study_id,local_case_id,study_groups_json,eligibility_status,selected_measurements_json,steiner_reference FROM cases WHERE case_id=?",(storage_id,)).fetchone()
        super().load_case_from_database(storage_id)
        if row:
            workflow,study_id,local,groups,elig,measures,ref=row;self.workflow_type=workflow or "individual";self.current_study_id=study_id;self.current_study=self._get_study(study_id) if study_id else None;self.case_group_values=_loads(groups,{});self.selected_measure_keys=[k for k in _loads(measures,[]) if k in MEASUREMENTS] or ([k for k in self.current_study.get("measurements",[]) if k in MEASUREMENTS] if self.current_study else []);self.steiner_reference=ref or (self.current_study.get("steiner_reference") if self.current_study else "none") or "none";self.manual_eligibility="yes" if elig=="eligible" else ("no" if elig=="not_eligible" else "pending");self.case_id.set(local or storage_id);self._loaded_case_id=storage_id;self._apply_selected_measurements();self._update_compact_context()

    def new_case(self):
        workflow,study_id,study,keys,ref=self.workflow_type,self.current_study_id,self.current_study,list(self.selected_measure_keys),self.steiner_reference;result=super().new_case();self.workflow_type=workflow;self.current_study_id=study_id;self.current_study=study;self.selected_measure_keys=keys;self.steiner_reference=ref;self.case_group_values={};self.manual_eligibility="pending";self.case_id.set("");self._apply_selected_measurements();self._update_compact_context();return result

    def next_case(self):
        if self.original and not self._case_saved and not messagebox.askyesno("Nuevo caso","Hay cambios sin guardar. ¿Continuar sin guardarlos?"):return
        next_id=""
        if self.workflow_type=="research" and self.current_study_id:
            with sqlite3.connect(self._db_path) as con:rows=[r[0] for r in con.execute("SELECT local_case_id FROM cases WHERE research_study_id=?",(self.current_study_id,)).fetchall() if str(r[0] or "").isdigit()]
            next_id=str(max([int(x) for x in rows],default=0)+1)
        self.new_case();self.case_id.set(next_id);self.open_case_data(first_time=True)

    def open_database_browser(self):
        win=tk.Toplevel(self);win.title("YomCeph · Base de datos");sw,sh=self.winfo_screenwidth(),self.winfo_screenheight();w,h=min(1100,int(sw*.9)),min(720,int(sh*.85));win.geometry(f"{w}x{h}+{max(0,(sw-w)//2)}+{max(0,(sh-h)//2)}");fr=ttk.Frame(win,padding=10);fr.pack(fill=tk.BOTH,expand=True);cols=("case","study","age","sex","elig","updated");tree=ttk.Treeview(fr,columns=cols,show="headings")
        for col,title,width in [("case","Caso",100),("study","Investigación / tipo",300),("age","Edad",80),("sex","Sexo",130),("elig","Elegibilidad",120),("updated","Actualizado",190)]:tree.heading(col,text=title);tree.column(col,width=width,anchor="center" if col in ("case","age","elig") else "w")
        if self.workflow_type=="research" and self.current_study_id:q="""SELECT c.case_id,COALESCE(c.local_case_id,c.case_id),s.name,c.age_years,c.sex,c.eligibility_status,c.updated_at FROM cases c LEFT JOIN research_studies s ON s.study_id=c.research_study_id WHERE c.research_study_id=? ORDER BY CASE WHEN COALESCE(c.local_case_id,'') NOT GLOB '*[^0-9]*' THEN CAST(c.local_case_id AS INTEGER) ELSE 2147483647 END,c.local_case_id""";params=(self.current_study_id,)
        else:q="""SELECT c.case_id,COALESCE(c.local_case_id,c.case_id),'Caso individual',c.age_years,c.sex,c.eligibility_status,c.updated_at FROM cases c WHERE c.workflow_type='individual' ORDER BY c.updated_at DESC""";params=()
        with sqlite3.connect(self._db_path) as con:rows=con.execute(q,params).fetchall()
        for storage,local,study,age,sex,elig,upd in rows:tree.insert("",tk.END,iid=storage,values=(local,study,"" if age is None else age,sex or "",elig or "",upd))
        tree.pack(fill=tk.BOTH,expand=True);buttons=ttk.Frame(win,padding=(10,0,10,10));buttons.pack(fill=tk.X)
        def open_sel():
            sel=tree.selection()
            if sel:win.destroy();self.load_case_from_database(sel[0])
        ttk.Button(buttons,text="Abrir caso",command=open_sel,style="Purple.TButton").pack(side=tk.LEFT);ttk.Button(buttons,text="Exportar investigación",command=self.export_research_package,style="Gold.TButton").pack(side=tk.LEFT,padx=6)

    def _update_compact_context(self):
        if not hasattr(self,"workflow_badge"):return
        if self.workflow_type=="research" and self.current_study:self.workflow_badge.config(text="INVESTIGACIÓN",bg=ui.C["purple_dark"]);self.study_compact_label.config(text=self.current_study["name"])
        else:self.workflow_badge.config(text="CASO INDIVIDUAL",bg=ui.C["turquoise_dark"]);self.study_compact_label.config(text="No investigación")
        self._update_db_counter()

    def _update_db_counter(self):
        if not hasattr(self,"db_counter") or not self._db_path:return
        try:
            with sqlite3.connect(self._db_path) as con:
                if self.workflow_type=="research" and self.current_study_id:total=con.execute("SELECT COUNT(*) FROM cases WHERE research_study_id=?",(self.current_study_id,)).fetchone()[0];included=con.execute("SELECT COUNT(*) FROM cases WHERE research_study_id=? AND eligibility_status='eligible'",(self.current_study_id,)).fetchone()[0];target=(self.current_study or {}).get("target_n","?");text=f"{included}/{target} incluidos · {total} registrados"
                else:total=con.execute("SELECT COUNT(*) FROM cases WHERE workflow_type='individual'").fetchone()[0];text=f"{total} individuales"
            self.db_counter.config(text=text)
        except Exception:pass

    def _research_rows(self):
        if self.workflow_type!="research" or not self.current_study_id:return [],{}
        with sqlite3.connect(self._db_path) as con:
            cases=con.execute("""SELECT case_id,COALESCE(local_case_id,case_id),age_years,age_months,sex,sex_code,gender_identity,country,institution,clinic,radiograph_date,study_groups_json,eligibility_status,eligibility_reason,updated_at FROM cases WHERE research_study_id=? ORDER BY CASE WHEN COALESCE(local_case_id,'') NOT GLOB '*[^0-9]*' THEN CAST(local_case_id AS INTEGER) ELSE 2147483647 END,local_case_id""",(self.current_study_id,)).fetchall();mids=[r[0] for r in cases];measures={}
            if mids:
                placeholders=",".join("?" for _ in mids)
                for cid,name,value,diagnosis in con.execute(f"SELECT case_id,name,value,diagnosis FROM measurements WHERE case_id IN ({placeholders})",mids):measures.setdefault(cid,{})[name]=(value,diagnosis or "")
        return cases,measures

    def export_research_package(self):
        if self.workflow_type!="research" or not self.current_study:messagebox.showinfo("Exportar","Los paquetes estadísticos se generan únicamente para una investigación activa.");return
        folder=filedialog.askdirectory(title="Carpeta para exportar la investigación")
        if not folder:return
        cases,measure_data=self._research_rows();study=self.current_study;selected=[k for k in study.get("measurements",[]) if k in MEASUREMENTS];group_names=[f.get("name") for f in study.get("group_fields",[])];headers=["Caso","Edad_años","Edad_meses","Sexo","Codigo_sexo","Genero","Pais","Institucion","Clinica","Fecha_RX","Elegibilidad","Motivo_elegibilidad"]+group_names+selected;matrix=[]
        for row in cases:
            cid,local,years,months,sex,sexcode,gender,country,institution,clinic,rx,groups_raw,elig,reason,updated=row;groups=_loads(groups_raw,{});out=[local,"" if years is None else years,"" if months is None else months,sex or "",sexcode or "",gender or "",country or "",institution or "",clinic or "",rx or "",elig or "",reason or ""]+[groups.get(g,"") for g in group_names];m=measure_data.get(cid,{});out += [m.get(k,("",))[0] if k in m else "" for k in selected];matrix.append(out)
        base=_safe_id(study["name"]);csv_path=Path(folder)/f"{base}_datos_SPSS.csv"
        with open(csv_path,"w",newline="",encoding="utf-8-sig") as f:writer=csv.writer(f);writer.writerow(headers);writer.writerows(matrix)
        self._write_spss_syntax(Path(folder)/f"{base}_importar_SPSS.sps",csv_path,headers,selected);xlsx_path=Path(folder)/f"{base}_YomCeph.xlsx";self._write_excel(xlsx_path,headers,matrix,selected,group_names);self.status.config(text=f"✓ Investigación exportada: {folder}");messagebox.showinfo("Exportar",f"Archivos creados:\n{xlsx_path.name}\n{csv_path.name}\n{base}_importar_SPSS.sps")

    def _write_excel(self,path,headers,matrix,selected,group_names):
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
        except Exception:messagebox.showwarning("Excel","Falta openpyxl; se conservarán CSV y sintaxis SPSS.");return
        wb=Workbook();ws=wb.active;ws.title="Datos";ws.append(headers)
        for row in matrix:ws.append(row)
        ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
        for i,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=min(28,max(10,len(str(h))+2))
        desc=wb.create_sheet("Descriptivos");desc.append(["Variable","N","Media","DE","Mínimo","Máximo"]);idx={h:i for i,h in enumerate(headers)};included=[r for r in matrix if r[idx["Elegibilidad"]]=="eligible"]
        for key in selected:
            vals=[]
            for row in included:
                v=row[idx[key]]
                if isinstance(v,(int,float)):vals.append(float(v))
            desc.append([key,len(vals),statistics.mean(vals) if vals else "",statistics.stdev(vals) if len(vals)>1 else "",min(vals) if vals else "",max(vals) if vals else ""])
        freq=wb.create_sheet("Frecuencias");freq.append(["Variable","Categoría","N","Porcentaje"])
        for h in ["Sexo","Elegibilidad"]+group_names:
            counts={}
            for row in matrix:
                value=str(row[idx[h]] or "(vacío)");counts[value]=counts.get(value,0)+1
            total=sum(counts.values()) or 1
            for value,n in sorted(counts.items()):freq.append([h,value,n,n*100/total])
        prev=wb.create_sheet("Prevalencias");prev.append(["Variable","Categoría","N","Porcentaje","Nota"])
        for key in selected:
            categories={}
            for row in included:
                v=row[idx[key]]
                if isinstance(v,(int,float)):
                    cat=self._classification_for(key,float(v))
                    if cat:categories[cat]=categories.get(cat,0)+1
            total=sum(categories.values())
            for cat,n in sorted(categories.items()):prev.append([key,cat,n,(n*100/total if total else ""),"Sólo categorías con regla explícita del protocolo/referencia"])
        dic=wb.create_sheet("Diccionario");dic.append(["Variable","Etiqueta","Unidad","Análisis","Landmarks"])
        for key in selected:
            s=MEASUREMENTS[key];dic.append([key,s.get("label",key),s.get("unit",""),s.get("analysis",""),", ".join(s.get("required",[]))])
        prot=wb.create_sheet("Protocolo");st=self.current_study
        for k,v in [("Nombre",st["name"]),("Muestra planeada",st["target_n"]),("Permite aumentar",bool(st["allow_target_increase"])),("Edad mínima",st.get("age_min")),("Edad máxima",st.get("age_max")),("País",st.get("country")),("Institución",st.get("institution")),("Inclusión",st.get("inclusion_criteria")),("Exclusión",st.get("exclusion_criteria")),("Referencia Steiner",st.get("steiner_reference")),("Versión protocolo",st.get("protocol_version"))]:prot.append([k,"" if v is None else v])
        prot.append(["Análisis",", ".join(st.get("analyses",[]))]);prot.append(["Mediciones",", ".join(selected)]);wb.save(path)

    def _write_spss_syntax(self,path,csv_path,headers,selected):
        vars_=[_clean_spss_name(h) for h in headers];lines=["* YomCeph v0.12 - importación reproducible.",f"GET DATA /TYPE=TXT /FILE='{str(csv_path).replace(chr(39),chr(39)*2)}' /ENCODING='UTF8'","/DELCASE=LINE /DELIMITERS=\",\" /QUALIFIER='\"' /ARRANGEMENT=DELIMITED /FIRSTCASE=2","/VARIABLES="];numeric=set(selected)|{"Edad_años","Edad_meses"}
        for original,var in zip(headers,vars_):lines.append(f" {var} F12.4" if original in numeric else f" {var} A120")
        lines[-1]+=".";numvars=[_clean_spss_name(k) for k in selected]
        if numvars:lines.append("DESCRIPTIVES VARIABLES="+" ".join(numvars)+" /STATISTICS=MEAN STDDEV MIN MAX.")
        lines.append("FREQUENCIES VARIABLES="+_clean_spss_name("Elegibilidad")+" "+_clean_spss_name("Sexo")+".");lines.append("EXECUTE.");Path(path).write_text("\n".join(lines),encoding="utf-8-sig")


if __name__ == "__main__":
    ui.enable_dpi_awareness()
    app=YomCephV120()
    app.mainloop()
