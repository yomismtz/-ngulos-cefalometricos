import json
import sqlite3
import tkinter as tk
from datetime import datetime
from tkinter import messagebox, simpledialog

import yomceph_desktop_hidpi as ui
import yomceph_desktop_v120 as v120
from yomceph_desktop_v120_release import YomCephV120Release, _spss_base_name
from yomceph_v120_geometry import legacy_active_angles

APP_VERSION = "0.12.0"
MAX_RESEARCH_CASES = 1000


class YomCephV120ReleaseFinal(YomCephV120Release):
    """Entrada final de la candidata pública v0.12.

    La cadena histórica todavía contiene ``closest_supplement`` para abrir
    proyectos creados con versiones anteriores. Esta clase reemplaza, antes de
    mostrar o persistir resultados de v0.12, todas las mediciones activas que lo
    usaban por construcciones geométricas explícitas e independientes de normas.
    """

    def _init_database(self):
        super()._init_database()
        with sqlite3.connect(self._db_path) as con:
            con.execute(
                """CREATE TABLE IF NOT EXISTS research_protocol_history (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    study_id TEXT NOT NULL,
                    action TEXT NOT NULL,
                    changed_at TEXT NOT NULL,
                    protocol_version INTEGER NOT NULL,
                    details_json TEXT NOT NULL DEFAULT '{}'
                )"""
            )
            con.execute(
                "CREATE INDEX IF NOT EXISTS idx_protocol_history_study ON research_protocol_history(study_id,id)"
            )
            con.execute(
                """UPDATE research_studies
                   SET steiner_reference='none'
                   WHERE steiner_reference='legacy_imported'
                     AND study_id NOT GLOB 'investigacion_importada*'"""
            )
            con.execute(
                """UPDATE cases
                   SET steiner_reference='none'
                   WHERE steiner_reference='legacy_imported'
                     AND (research_study_id IS NULL
                          OR research_study_id NOT GLOB 'investigacion_importada*')"""
            )
            con.commit()

    def _record_protocol_history(self, study_id, action, details=None, version=None):
        if not study_id:
            return
        now = datetime.now().isoformat(timespec="seconds")
        with sqlite3.connect(self._db_path) as con:
            con.execute(
                """CREATE TABLE IF NOT EXISTS research_protocol_history (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    study_id TEXT NOT NULL,
                    action TEXT NOT NULL,
                    changed_at TEXT NOT NULL,
                    protocol_version INTEGER NOT NULL,
                    details_json TEXT NOT NULL DEFAULT '{}'
                )"""
            )
            if version is None:
                row = con.execute(
                    "SELECT protocol_version FROM research_studies WHERE study_id=?",
                    (study_id,),
                ).fetchone()
                version = int(row[0]) if row and row[0] is not None else 1
            con.execute(
                """INSERT INTO research_protocol_history(
                    study_id,action,changed_at,protocol_version,details_json
                ) VALUES(?,?,?,?,?)""",
                (
                    study_id,
                    action,
                    now,
                    int(version),
                    json.dumps(details or {}, ensure_ascii=False, sort_keys=True),
                ),
            )
            con.commit()

    def _install_database_bar(self):
        super()._install_database_bar()
        try:
            for child in self.universal_wrapper.winfo_children():
                if str(child.cget("text")) != "Más ▾":
                    continue
                menu = self.nametowidget(str(child.cget("menu")))
                menu.add_separator()
                menu.add_checkbutton(
                    label="Alto contraste",
                    variable=self.high_contrast_var,
                    command=self._apply_high_contrast,
                )
                break
        except Exception:
            pass

    def calculate_values(self):
        values = super().calculate_values()
        values.update(legacy_active_angles(self.points))
        return values

    def _save_study(self, data, existing_id=None):
        # Evita colisiones humanas entre variables de grupo y mediciones/columnas
        # derivadas. El padre también normaliza duplicados y nombres reservados.
        reserved = set()
        for key in v120.MEASUREMENTS:
            reserved.add(_spss_base_name(key).casefold())
            reserved.add(_spss_base_name(key + "__categoria").casefold())
        cleaned = dict(data)
        fields = []
        used = set(reserved)
        adjusted = False
        for field in data.get("group_fields", []):
            original = str(field.get("name", "Grupo")).strip() or "Grupo"
            name = original
            normalized = _spss_base_name(name).casefold()
            suffix = 2
            if normalized in used:
                name = original + " grupo"
                normalized = _spss_base_name(name).casefold()
                adjusted = True
            while normalized in used:
                name = f"{original} grupo {suffix}"
                normalized = _spss_base_name(name).casefold()
                suffix += 1
                adjusted = True
            used.add(normalized)
            fields.append({"name": name, "options": list(field.get("options", []))})
        cleaned["group_fields"] = fields
        if adjusted:
            messagebox.showwarning(
                "Variables del estudio",
                "Una variable de grupo coincidía con el nombre de una medición. "
                "YomCeph la renombró para mantener columnas inequívocas en Excel y SPSS."
            )

        study_id = super()._save_study(cleaned, existing_id)
        now = datetime.now().isoformat(timespec="seconds")
        with sqlite3.connect(self._db_path) as con:
            if existing_id:
                con.execute(
                    """UPDATE research_studies
                       SET protocol_version=protocol_version+1, updated_at=?
                       WHERE study_id=?""",
                    (now, study_id),
                )
            row = con.execute(
                "SELECT protocol_version FROM research_studies WHERE study_id=?",
                (study_id,),
            ).fetchone()
            version = int(row[0]) if row and row[0] is not None else 1
            con.commit()

        snapshot = {
            "name": cleaned.get("name", ""),
            "target_n": cleaned.get("target_n"),
            "allow_target_increase": bool(cleaned.get("allow_target_increase")),
            "age_min": cleaned.get("age_min"),
            "age_max": cleaned.get("age_max"),
            "country": cleaned.get("country", ""),
            "institution": cleaned.get("institution", ""),
            "group_fields": cleaned.get("group_fields", []),
            "analyses": cleaned.get("analyses", []),
            "measurements": cleaned.get("measurements", []),
            "steiner_reference": cleaned.get("steiner_reference", "none"),
        }
        self._record_protocol_history(
            study_id,
            "edited" if existing_id else "created",
            snapshot,
            version=version,
        )
        return study_id

    def _ensure_research_capacity(self, local_case_id):
        """Evita rebasar silenciosamente la muestra planeada."""
        if self.workflow_type != "research" or not self.current_study_id or not self.current_study:
            return True
        status, _reason = self._evaluate_eligibility()
        if status != "eligible":
            return True

        storage_id = self._storage_case_id(local_case_id)
        with sqlite3.connect(self._db_path) as con:
            already_included = con.execute(
                """SELECT 1 FROM cases
                   WHERE case_id=? AND research_study_id=? AND eligibility_status='eligible'""",
                (storage_id, self.current_study_id),
            ).fetchone() is not None
            if already_included:
                return True
            included = con.execute(
                """SELECT COUNT(*) FROM cases
                   WHERE research_study_id=? AND eligibility_status='eligible'""",
                (self.current_study_id,),
            ).fetchone()[0]

        target = int(self.current_study.get("target_n") or 0)
        if included < target:
            return True
        if included >= MAX_RESEARCH_CASES or target >= MAX_RESEARCH_CASES:
            messagebox.showwarning(
                "Capacidad de investigación",
                f"YomCeph v0.12 admite hasta {MAX_RESEARCH_CASES} casos planeados por investigación. "
                "No se puede ampliar más esta muestra."
            )
            return False
        if not bool(self.current_study.get("allow_target_increase")):
            messagebox.showwarning(
                "Muestra completa",
                f"La investigación ya alcanzó la muestra planeada de {target} casos incluidos.\n\n"
                "El protocolo no permite ampliarla. Este caso no se guardará como incluido."
            )
            return False

        new_target = simpledialog.askinteger(
            "Ampliar muestra",
            f"La investigación ya alcanzó {included}/{target} casos incluidos.\n\n"
            "Para incluir este paciente, registre primero el nuevo tamaño planeado de muestra:",
            minvalue=included + 1,
            maxvalue=MAX_RESEARCH_CASES,
            parent=self,
        )
        if not new_target:
            return False
        now = datetime.now().isoformat(timespec="seconds")
        with sqlite3.connect(self._db_path) as con:
            con.execute(
                """UPDATE research_studies
                   SET target_n=?, protocol_version=protocol_version+1, updated_at=?
                   WHERE study_id=?""",
                (new_target, now, self.current_study_id),
            )
            version = con.execute(
                "SELECT protocol_version FROM research_studies WHERE study_id=?",
                (self.current_study_id,),
            ).fetchone()[0]
            con.commit()
        self._record_protocol_history(
            self.current_study_id,
            "sample_increased",
            {
                "old_target": target,
                "new_target": new_target,
                "included_at_change": included,
            },
            version=version,
        )
        self.current_study = self._get_study(self.current_study_id)
        self.study_target_var.set(str(new_target))
        self._update_db_counter()
        messagebox.showinfo(
            "Muestra ampliada",
            f"El tamaño planeado quedó registrado como {new_target}. "
            "La versión del protocolo se incrementó y el cambio quedó en el historial."
        )
        return True

    def _validate_included_case_completeness(self):
        if self.workflow_type != "research" or not self.current_study:
            return True
        status, _reason = self._evaluate_eligibility()
        if status != "eligible":
            return True
        if not self.sex_code_var.get().strip():
            messagebox.showwarning(
                "Dato pendiente",
                "Antes de incluir el caso seleccione el sexo registrado. Si no está disponible, "
                "utilice una categoría explícita como “No registrado”."
            )
            return False
        missing = [
            field.get("name", "Grupo")
            for field in self.current_study.get("group_fields", [])
            if not self.case_group_values.get(field.get("name", ""), "").strip()
        ]
        if missing:
            messagebox.showwarning(
                "Grupo pendiente",
                "Antes de incluir el caso complete las variables del protocolo: "
                + ", ".join(missing)
                + ".\n\nUse Datos del caso. Si una variable puede faltar, añada al protocolo "
                  "una opción explícita como “No registrado” o “No aplica”."
            )
            return False
        return True

    def _write_excel_release(
        self, path, headers, matrix, selected, group_names, category_headers, spss_map
    ):
        super()._write_excel_release(
            path, headers, matrix, selected, group_names, category_headers, spss_map
        )
        from openpyxl import load_workbook
        from openpyxl.styles import Font

        with sqlite3.connect(self._db_path) as con:
            rows = con.execute(
                """SELECT changed_at,protocol_version,action,details_json
                   FROM research_protocol_history
                   WHERE study_id=? ORDER BY id""",
                (self.current_study_id,),
            ).fetchall()

        wb = load_workbook(path)
        if "Historial protocolo" in wb.sheetnames:
            del wb["Historial protocolo"]
        ws = wb.create_sheet("Historial protocolo")
        ws.append(["Fecha", "Versión", "Acción", "Detalles"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for changed_at, version, action, details_json in rows:
            try:
                details = json.loads(details_json or "{}")
                details_text = json.dumps(details, ensure_ascii=False, sort_keys=True)
            except Exception:
                details_text = details_json or ""
            ws.append([changed_at, version, action, details_text])
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 10
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 90
        wb.save(path)

    def save_case_to_database(self):
        local = self.case_id.get().strip()
        if local and not self._validate_included_case_completeness():
            return
        if local and not self._ensure_research_capacity(local):
            return
        return super().save_case_to_database()


if __name__ == "__main__":
    ui.enable_dpi_awareness()
    app = YomCephV120ReleaseFinal()
    app.mainloop()
